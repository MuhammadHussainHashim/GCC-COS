"""Evidence and coverage audit. A supplier/mirror review claim is not an independent check."""
from collections import Counter, defaultdict
from decimal import Decimal, InvalidOperation
import csv
import hashlib
import json
from pathlib import Path
from urllib.parse import urlsplit
import pandas as pd
from .ingest import load_registry
from .model import is_product
from .products import ilcd_records, year_of, boavizta
from .storage import records, now

TARGETS = ['cement', 'steel', 'aluminum', 'glass', 'vehicles', 'electronics', 'appliances', 'furniture', 'clothing']


def key(value):
    return ' '.join(str(value or '').casefold().split())


def signature(r):
    try:
        value = str(Decimal(str(r.get('original_value'))).normalize())
    except InvalidOperation:
        value = str(r.get('original_value'))
    return tuple(key(r.get(k)) for k in ('manufacturer', 'product_name', 'publication_date', 'geography',
                                         'lifecycle_scope', 'declared_quantity', 'declared_unit', 'original_unit')) + (value,)


def independent_class(record, trace_ok):
    # There is no independently executed primary-document verification workflow yet.
    # Neither import metadata, a URL nor manual acceptance is sufficient to establish one.
    return 'published_third_party_data' if trace_ok else 'unverified_record'


def _document(entry, source):
    result = {'archive_status': 'missing_archive', 'license_review': 'not_established', 'lookup': {}}
    if not entry:
        return result
    sha, sid, url, title, retrieved, filename, extracted = entry
    path = Path(filename)
    if not path.is_file():
        return result
    if hashlib.sha256(path.read_bytes()).hexdigest() != sha:
        result['archive_status'] = 'archive_checksum_mismatch'
        return result
    if not source or source.get('sha256') != sha or source.get('url') != url:
        result['archive_status'] = 'registry_identity_or_checksum_mismatch'
        return result
    result['archive_status'] = 'checksum_matches_pinned_dataset'
    try:
        if source.get('adapter') == 'boavizta':
            license_file = Path(source.get('license_evidence_file', ''))
            if not license_file.is_file():
                return result
            content = license_file.read_bytes()
            if hashlib.sha256(content).hexdigest() != source.get('license_evidence_sha256'):
                result['license_review'] = 'license_evidence_checksum_mismatch'
                return result
            terms = content.decode('utf-8')
            if 'Open Database License' not in terms or 'Database Contents License' not in terms:
                return result
            result['license_review'] = 'ODbL_DbCL_declared_in_pinned_README'
            with path.open(encoding='utf-8-sig', newline='') as handle:
                result['lookup'] = dict(enumerate(csv.DictReader(handle), 2))
        elif source.get('adapter') == 'ilcd_epd':
            doc = json.loads(path.read_text(encoding='utf-8'))
            # This adapter validates licenseType/access restrictions inside the document.
            parsed = list(ilcd_records(doc, source))
            result['license_review'] = 'embedded_all_uses_permission_no_access_restrictions'
            result['lookup'] = {r['lifecycle_module']: r for r in parsed}
    except (ValueError, KeyError, OSError, TypeError) as exc:
        result['license_review'] = 'license_or_schema_review_failed:' + type(exc).__name__
    return result


def _matches(record, document, source):
    adapter = source.get('adapter')
    expected = None
    if adapter == 'boavizta':
        raw = document['lookup'].get((record.get('evidence') or {}).get('row'))
        if raw:
            expected = boavizta(raw, source)
    elif adapter == 'ilcd_epd':
        expected = document['lookup'].get(record.get('lifecycle_module'))
    if not expected:
        return False
    for field, value in expected.items():
        if field == 'original_value':
            try:
                if Decimal(str(record.get(field))) != Decimal(str(value)):
                    return False
            except InvalidOperation:
                return False
        elif record.get(field) != value:
            return False
    return True


def build_audit(db, registry='config/sources.json'):
    sources = {s['id']: s for s in load_registry(registry)}
    products = [r for r in records(db) if is_product(r)]
    documents = {(e[0], e[1]): e for e in db.execute('SELECT * FROM documents')}
    document_cache = {}
    for r in products:
        k = (r.get('document_sha256'), r.get('source_id'))
        if k not in document_cache:
            document_cache[k] = _document(documents.get(k), sources.get(r.get('source_id')))
    groups, repeated, exact_content = defaultdict(list), defaultdict(list), defaultdict(list)
    for r in products:
        groups[(key(r.get('manufacturer')), key(r.get('product_name')))].append(r)
        repeated[signature(r)].append(r)
        evidence = r.get('evidence') or {}
        content = evidence.get('raw', evidence)
        exact_content[(r.get('source_id'), r.get('document_sha256'), json.dumps(content,sort_keys=True))].append(r)
    duplicate_groups = []
    for group in groups.values():
        if len(group) > 1:
            scopes = {r.get('lifecycle_scope') for r in group}
            duplicate_groups.append({'manufacturer': group[0].get('manufacturer'), 'product_name': group[0].get('product_name'),
                                     'record_ids': [r['id'] for r in group], 'observations': len(group),
                                     'reason': 'different_lifecycle_modules' if len(scopes) > 1 else 'same_name_multiple_observations_review_required',
                                     'automatically_merged': False,
                                     'original_publisher_references': sorted({r.get('primary_source_url') or '' for r in group}),
                                     'lifecycle_boundaries': sorted({r.get('lifecycle_scope') or '' for r in group}),
                                     'footprint_values': [r.get('original_value') for r in group]})
    candidate_ids = {r['id'] for group in repeated.values() if len(group) > 1 for r in group}
    observations = []
    source_rows = []
    for r in products:
        document = document_cache[(r.get('document_sha256'), r.get('source_id'))]
        source = sources.get(r.get('source_id'), {})
        matches = _matches(r, document, source)
        license_ok = document['license_review'] in {'ODbL_DbCL_declared_in_pinned_README', 'embedded_all_uses_permission_no_access_restrictions'}
        attribution_ok = bool(r.get('manufacturer') and r.get('attribution') == source.get('attribution') and r.get('attribution') and r.get('license') == source.get('license') and r.get('license_url') == source.get('license_url') and r.get('source_url') == source.get('url') and r.get('document_title') == source.get('title'))
        trace_ok = matches and license_ok and attribution_ok and document['archive_status'] == 'checksum_matches_pinned_dataset'
        warnings = list(r.get('issues') or [])
        if not matches: warnings.append('record_does_not_match_archived_source_or_unsupported_adapter')
        if not license_ok: warnings.append('reuse_permission_evidence_requires_review')
        if not attribution_ok: warnings.append('attribution_or_license_metadata_requires_review')
        if r['id'] in candidate_ids: warnings.append('possible_repeated_footprint_signature')
        if r.get('verification_status') == 'independently_crosschecked': warnings.append('unsupported_independent_verification_claim')
        observations.append({**r, 'model_or_variant': r.get('product_name'),
                             'model_note': 'Original product designation; no separate variant inferred',
                             'publication_year': year_of(str(r.get('publication_date') or '')),
                             'reference_year': r.get('year'), 'functional_unit': f"{r.get('declared_quantity') or '?'} {r.get('declared_unit') or '?'}",
                             'audit_verification_class': independent_class(r, trace_ok),
                             'archive_status': document['archive_status'], 'source_row_matches_archive': matches,
                             'license_review': document['license_review'], 'primary_publisher_crosscheck': 'not_performed',
                             'possible_duplicate_signature': r['id'] in candidate_ids, 'audit_warnings': warnings})
    for sid in sorted({r.get('source_id') for r in products}, key=str):
        subset = [r for r in observations if r['source_id'] == sid]
        s = sources.get(sid, {})
        source_rows.append({'source_id': sid, 'observations': len(subset), 'source_url': s.get('url'),
                            'license': s.get('license'), 'license_url': s.get('license_url'),
                            'license_evidence_url': s.get('license_evidence_url'), 'attribution': s.get('attribution'),
                            'license_review': sorted({r['license_review'] for r in subset}),
                            'archive_status': sorted({r['archive_status'] for r in subset}),
                            'publisher_crosscheck': 'not_performed',
                            'declaration_review_claims': dict(Counter(r.get('verification_status') for r in subset)),
                            'applicable_obligations': 'Attribution and applicable ODbL share-alike / DbCL terms; manufacturer PDFs not licensed by this collection' if s.get('adapter') == 'boavizta' else 'Preserve dataset owner attribution, embedded terms and all access restrictions; software license is not a data license'})
    coverage = []
    for target in TARGETS:
        if target == 'electronics': subset = [r for r in products if r.get('category') == 'electronics']
        elif target in ('aluminum', 'glass', 'steel'): subset = [r for r in products if key(r.get('material')) == target]
        elif target == 'cement': subset = [r for r in products if 'cement' in key(r.get('material'))]
        else: subset = [r for r in products if r.get('category') == target]
        coverage.append({'category': target, 'observations': len(subset),
                         'distinct_manufacturer_product_names': len({(key(r.get('manufacturer')), key(r.get('product_name'))) for r in subset}),
                         'status': 'collected_with_review_limitations' if subset else 'no_records_collected',
                         'limitation': 'Cement-based foam/boards are not plain cement; concrete is reported separately' if target == 'cement' else 'Electronics includes limited home/IoT/printers; not general appliance coverage' if target == 'electronics' else 'No coverage inferred from registry placeholders'})
    concrete = [r for r in products if r.get('material') == 'concrete']
    coverage.append({'category': 'concrete', 'observations': len(concrete), 'distinct_manufacturer_product_names': len({(r.get('manufacturer'),r.get('product_name')) for r in concrete}), 'status': 'collected_with_review_limitations' if concrete else 'no_records_collected','limitation': 'Separate lifecycle modules are not separate products'})
    for material in ('plain cement','ceramics'):
        subset=[r for r in products if r.get('material') == ('cement' if material=='plain cement' else 'ceramics')]
        coverage.append({'category':material,'observations':len(subset),'distinct_manufacturer_product_names':len({(r.get('manufacturer'),r.get('product_name')) for r in subset}),'status':'collected_with_review_limitations' if subset else 'no_records_collected','limitation':'No extrapolation from cement-based boards/foam or other materials'})
    summary = {'generated_at': now(), 'product_observations': len(products), 'distinct_manufacturer_product_names': len(groups),
               'unique_SKUs': 'not_established', 'source_datasets': len(source_rows),
               'verification_classes': dict(Counter(r['audit_verification_class'] for r in observations)),
               'independently_verified': 0, 'primary_document_crosschecks_performed': 0,
               'source_reported_review_statuses': dict(Counter(str(r.get('verification_status')) for r in products)),
               'primary_publisher_references_present': sum(bool(r.get('primary_source_url')) for r in products),
               'same_name_groups': len(duplicate_groups), 'repeated_signature_groups': sum(len(v)>1 for v in repeated.values()),
               'repeated_signature_records': len(candidate_ids), 'automatically_merged_records': 0,
               'exact_repeated_source_content_groups': sum(len(group)>1 for group in exact_content.values()),
               'exact_repeated_source_content_records': sum(len(group) for group in exact_content.values() if len(group)>1),
               'traceable_source_rows': sum(r['source_row_matches_archive'] for r in observations),
               'declared_units': dict(Counter(str(r.get('declared_unit')) for r in products)),
               'lifecycle_boundaries': dict(Counter(str(r.get('lifecycle_scope')) for r in products)),
               'licenses': dict(Counter(str(r.get('license')) for r in products)),
               'warnings': dict(Counter(i for r in observations for i in r['audit_warnings'])),
               'records_requiring_human_review': sum(bool(r['audit_warnings']) for r in observations)}
    return {'summary': summary, 'observations': observations, 'sources': source_rows,
            'duplicates': duplicate_groups, 'coverage': coverage}


def safe_frame(rows):
    rows = [dict(r) for r in rows]
    for r in rows:
        for k, v in r.items():
            if isinstance(v, (list, dict)): r[k] = json.dumps(v, ensure_ascii=False)
            elif isinstance(v, str) and v.startswith(('=', '+', '-', '@', '\t', '\r')): r[k] = "'" + v
    return pd.DataFrame(rows)


def write_audit(db, registry, output_dir, workbook=None):
    target = Path(workbook) if workbook else Path(output_dir)/'products-audited.xlsx'
    origin = db.execute('PRAGMA database_list').fetchone()[2]
    if target.suffix.lower() != '.xlsx' or (origin and Path(origin).resolve() == target.resolve()):
        raise ValueError('Choose a separate .xlsx audit workbook; the active database cannot be overwritten')
    report = build_audit(db, registry)
    destination = Path(output_dir); destination.mkdir(parents=True, exist_ok=True)
    (destination/'summary.json').write_text(json.dumps(report['summary'], indent=2, ensure_ascii=False)+'\n', encoding='utf-8')
    for name in ('observations', 'sources', 'duplicates', 'coverage'):
        safe_frame(report[name]).to_csv(destination/(name+'.csv'), index=False, encoding='utf-8-sig')
    target = Path(workbook) if workbook else destination/'products-audited.xlsx'
    target.parent.mkdir(parents=True, exist_ok=True)
    # Reuse the normal export for original records, license notices and existing audits.
    from .export import excel_bytes
    from openpyxl import load_workbook
    import io
    book = load_workbook(io.BytesIO(excel_bytes(db, products_only=True)))
    for name, rows in [('Product audit', report['observations']), ('Source verification', report['sources']),
                       ('Coverage', report['coverage']), ('Duplicate review', report['duplicates']),
                       ('Audit summary', [{'key':k,'value':v} for k,v in report['summary'].items()])]:
        frame = safe_frame(rows)
        if any(isinstance(v,str) and len(v)>32767 for v in frame.to_numpy().flat):
            raise ValueError('Audit evidence exceeds Excel cell limit; inspect CSV/SQLite instead')
        sheet=book.create_sheet(name)
        if len(frame.columns): sheet.append(list(frame.columns))
        for values in frame.itertuples(index=False, name=None):
            sheet.append([None if pd.isna(v) else v for v in values])
        sheet.freeze_panes='A2'; sheet.auto_filter.ref=sheet.dimensions
    book.save(target)
    return report['summary']
